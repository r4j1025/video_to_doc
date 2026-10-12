"""Build the Word document with python-docx."""
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from llm_service import fmt_time


def _nearest_frame(frames, seconds):
    if not frames:
        return None
    return min(frames, key=lambda f: abs(f["time"] - seconds))


def build_document(content, frames, segments, output_path, source_name, include_transcript=True):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    # Title block
    doc.add_heading(content.get("title") or "Video Documentation", level=0)
    meta = doc.add_paragraph(f"Source video: {source_name}  |  Generated: {date.today():%d %b %Y}")
    meta.runs[0].italic = True

    # Summary
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(content.get("summary", ""))

    topics = content.get("key_topics") or []
    if topics:
        doc.add_heading("Key Topics", level=2)
        for t in topics:
            doc.add_paragraph(str(t), style="List Bullet")

    # Sections with screenshots
    for i, sec in enumerate(content.get("sections") or [], start=1):
        start = float(sec.get("start_seconds") or 0)
        doc.add_heading(f"{i}. {sec.get('heading', 'Section')}  ({fmt_time(start)})", level=1)

        if sec.get("explanation"):
            doc.add_paragraph(sec["explanation"])

        frame = _nearest_frame(frames, start)
        if frame:
            doc.add_picture(frame["path"], width=Inches(6))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph(f"Screenshot at {fmt_time(frame['time'])}")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True
            cap.runs[0].font.size = Pt(9)

        # numbered by hand: python-docx's "List Number" style does not restart per section
        for n, step in enumerate(sec.get("steps") or [], start=1):
            doc.add_paragraph(f"{n}. {step}")

    # Transcript appendix
    if include_transcript and segments:
        doc.add_page_break()
        doc.add_heading("Appendix: Full Transcript", level=1)
        for seg in segments:
            p = doc.add_paragraph()
            r = p.add_run(f"[{fmt_time(seg['start'])}] ")
            r.bold = True
            p.add_run(seg["text"])

    doc.save(str(output_path))
    return output_path
